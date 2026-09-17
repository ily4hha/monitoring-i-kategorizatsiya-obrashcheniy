"""Bounded Excel preflight followed by sequential sheet reading.

XLSX dimensions are untrusted: inspect XML events without building a tree, then
pass explicit row/column bounds to openpyxl. BIFF5/8 records receive the same
checks before xlrd is allowed to allocate a sheet.
"""
from __future__ import annotations

import posixpath
import re
import struct
import zipfile
from contextlib import closing, contextmanager
from dataclasses import dataclass
from io import BytesIO, StringIO
from xml.etree import ElementTree
from xml.parsers import expat

import openpyxl
import xlrd
from xlrd.compdoc import CompDoc


class DatasetError(ValueError):
    pass


@dataclass(frozen=True)
class ExcelLimits:
    sheets: int
    rows: int
    columns: int
    cells: int
    sparse_ratio: int
    book_rows: int
    book_cells: int
    uncompressed_bytes: int
    metadata_bytes: int = 8 * 1024 * 1024
    archive_entries: int = 1000

    def check_shape(self, name: str, rows: int, columns: int) -> None:
        data_rows = max(0, rows - 1)
        if columns > self.columns or data_rows > self.rows or data_rows * columns > self.cells:
            raise DatasetError(f"Лист «{name}» превышает ограничение по размеру")


@dataclass
class SheetBounds:
    name: str
    rows: int = 0
    columns: int = 0
    nonempty: int = 0
    physical_cells: int = 0

    def grow(self, row: int, column: int, limits: ExcelLimits) -> None:
        if row < 1 or column < 0:
            raise DatasetError("Некорректные координаты ячейки")
        self.rows = max(self.rows, row)
        self.columns = max(self.columns, column)
        limits.check_shape(self.name, self.rows, self.columns)

    def check_sparse(self, limits: ExcelLimits) -> None:
        if self.nonempty and max(0, self.rows - 1) * self.columns > self.nonempty * limits.sparse_ratio:
            raise DatasetError(f"Лист «{self.name}» слишком разрежен")


class WorkbookBudget:
    def __init__(self, limits: ExcelLimits):
        self.limits = limits
        self.rows = self.cells = 0

    def add(self, sheet: SheetBounds) -> None:
        sheet.check_sparse(self.limits)
        rows = max(0, sheet.rows - 1)
        self.rows += rows
        self.cells += rows * sheet.columns
        if self.rows > self.limits.book_rows or self.cells > self.limits.book_cells:
            raise DatasetError("Книга превышает общий бюджет строк или ячеек")


def _coordinate(value: str) -> tuple[int, int]:
    match = re.fullmatch(r"([A-Z]{1,3})([1-9][0-9]{0,6})", value)
    if not match:
        raise DatasetError("Некорректные координаты ячейки XLSX")
    column = 0
    for letter in match[1]:
        column = column * 26 + ord(letter) - ord("A") + 1
    return int(match[2]), column


def _scan_xml(source, limits: ExcelLimits, sheet: SheetBounds | None = None) -> None:
    """SAX-style parsing retains neither rows nor element trees, even for sparse XML."""
    parser = expat.ParserCreate(namespace_separator="}")
    depth = elements = row = column = cell_text = 0
    cell_row = 0
    value_depth = None
    has_value = False
    previous_row = 0

    def reject_dtd(*args):
        raise DatasetError("DTD и XML-сущности в XLSX недопустимы")

    def start(tag, attrs):
        nonlocal depth, elements, row, column, cell_row, cell_text, value_depth, has_value, previous_row
        depth += 1
        elements += 1
        if depth > 64 or elements > (limits.cells * 8 + 10000 if sheet else 250000):
            raise DatasetError("XLSX превышает бюджет XML-элементов")
        if sheet is None:
            return
        tag = tag.rsplit("}", 1)[-1]
        if tag == "dimension" and "ref" in attrs:
            end_row, end_column = _coordinate(attrs["ref"].split(":")[-1])
            limits.check_shape(sheet.name, end_row, end_column)
        elif tag == "row":
            row = int(attrs.get("r", row + 1))
            if row <= previous_row:
                raise DatasetError("Нарушен порядок строк XLSX")
            previous_row = row
            column = 0
            sheet.grow(row, 0, limits)
        elif tag == "c":
            cell_row, next_column = _coordinate(attrs["r"]) if "r" in attrs else (row, column + 1)
            if cell_row != row or next_column <= column:
                raise DatasetError("Нарушен порядок ячеек XLSX")
            column = next_column
            sheet.grow(cell_row, column, limits)
            sheet.physical_cells += 1
            if sheet.physical_cells > limits.cells + limits.columns:
                raise DatasetError("Лист превышает бюджет физических ячеек")
            cell_text = 0
            has_value = False
        elif tag in {"v", "t"}:
            value_depth = depth

    def chars(value):
        nonlocal cell_text, has_value
        if sheet is not None and value_depth is not None:
            cell_text += len(value)
            if cell_text > 32767:
                raise DatasetError("Текст ячейки XLSX превышает 32767 символов")
            has_value |= bool(value)

    def end(tag):
        nonlocal depth, value_depth
        if sheet is not None:
            if tag.rsplit("}", 1)[-1] == "c" and cell_row > 1 and has_value:
                sheet.nonempty += 1
            if depth == value_depth:
                value_depth = None
        depth -= 1

    parser.StartDoctypeDeclHandler = reject_dtd
    parser.EntityDeclHandler = reject_dtd
    parser.ExternalEntityRefHandler = reject_dtd
    parser.StartElementHandler = start
    parser.EndElementHandler = end
    parser.CharacterDataHandler = chars
    consumed = 0
    while chunk := source.read(64 * 1024):
        consumed += len(chunk)
        if consumed > limits.uncompressed_bytes:
            raise DatasetError("Распакованный XLSX превышает ограничение")
        parser.Parse(chunk, False)
    parser.Parse(b"", True)


def validate_xlsx_archive(content: bytes, limits: ExcelLimits) -> None:
    with zipfile.ZipFile(BytesIO(content)) as archive:
        infos = archive.infolist()
        if len(infos) > limits.archive_entries or len({info.filename for info in infos}) != len(infos):
            raise DatasetError("Недопустимое количество или повторение ZIP-компонентов")
        if sum(info.file_size for info in infos) > limits.uncompressed_bytes:
            raise DatasetError("Распакованный XLSX превышает ограничение 100 МБ")
        if any(info.flag_bits & 1 or info.compress_type not in {zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED} for info in infos):
            raise DatasetError("Неподдерживаемое сжатие или шифрование XLSX")


def _xlsx_bounds(content: bytes, limits: ExcelLimits) -> list[SheetBounds]:
    with zipfile.ZipFile(BytesIO(content)) as archive:
        metadata_used = 0
        checked = set()

        def metadata(name):
            nonlocal metadata_used
            metadata_used += archive.getinfo(name).file_size
            if metadata_used > limits.metadata_bytes:
                raise DatasetError("Метаданные XLSX превышают допустимый размер")
            data = archive.read(name)
            _scan_xml(BytesIO(data), limits)
            checked.add(name)
            return ElementTree.fromstring(data)

        workbook = metadata("xl/workbook.xml")
        relationships = metadata("xl/_rels/workbook.xml.rels")
        rels = {node.attrib["Id"]: node.attrib for node in relationships}
        nodes = workbook.findall("{*}sheets/{*}sheet")
        if len(nodes) > limits.sheets:
            raise DatasetError(f"Книга содержит больше {limits.sheets} листов")
        sheets = []
        paths = set()
        budget = WorkbookBudget(limits)
        for node in nodes:
            relation_id = node.attrib["{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"]
            relation = rels[relation_id]
            if relation.get("TargetMode") == "External" or not relation["Type"].endswith("/worksheet"):
                raise DatasetError("Поддерживаются только внутренние листы данных XLSX")
            target = relation["Target"]
            path = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join("xl", target))
            if not path.startswith("xl/") or path in paths or path in checked:
                raise DatasetError("Некорректная ссылка на лист XLSX")
            paths.add(path)
            sheet = SheetBounds(node.attrib["name"])
            with archive.open(path) as source:
                _scan_xml(source, limits, sheet)
            budget.add(sheet)
            sheets.append(sheet)
        # openpyxl eagerly loads styles and shared strings even in read-only mode.
        # Bound all remaining XML before allowing any such materialization.
        for info in archive.infolist():
            if info.filename not in paths | checked and info.filename.endswith((".xml", ".rels")):
                metadata(info.filename)
        return sheets


def _biff_records(stream: bytes, offset: int = 0):
    while offset + 4 <= len(stream):
        code, length = struct.unpack_from("<HH", stream, offset)
        end = offset + 4 + length
        if end > len(stream):
            raise DatasetError("Оборванная запись BIFF")
        yield code, stream[offset + 4:end]
        offset = end
        if code == 0x000A:  # EOF of this BIFF substream
            return
    raise DatasetError("Отсутствует конец потока BIFF")


def _xls_bounds(content: bytes, limits: ExcelLimits) -> list[SheetBounds]:
    document = CompDoc(content, logfile=StringIO())
    stream = document.get_named_stream("Workbook") or document.get_named_stream("Book")
    if not stream:
        raise DatasetError("В XLS отсутствует поток Workbook")
    records = _biff_records(stream)
    code, bof = next(records)
    if code != 0x0809 or struct.unpack_from("<H", bof)[0] not in {0x0500, 0x0600}:
        raise DatasetError("Поддерживается бинарный XLS формата BIFF5/BIFF8")
    version = struct.unpack_from("<H", bof)[0]
    offsets = []
    for code, data in records:
        if code == 0x0085:  # BOUNDSHEET
            if data[5] != 0:
                raise DatasetError("Поддерживаются только листы данных XLS")
            offsets.append(struct.unpack_from("<I", data)[0])
            if len(offsets) > limits.sheets:
                raise DatasetError(f"Книга содержит больше {limits.sheets} листов")
        elif code == 0x00FC:  # SST, checked before xlrd builds its string table
            unique = struct.unpack_from("<I", data, 4)[0]
            if unique > limits.book_cells + limits.sheets * limits.columns:
                raise DatasetError("XLS превышает бюджет общих строк")
    budget = WorkbookBudget(limits)
    sheets = []
    cell_codes = {0x0203, 0x027E, 0x00FD, 0x0204, 0x00D6, 0x0006, 0x0206, 0x0406, 0x0205, 0x0201}
    for index, offset in enumerate(offsets):
        sheet = SheetBounds(f"{index + 1}")
        for code, data in _biff_records(stream, offset):
            if code == 0x0200 and data:  # DIMENSIONS; still check actual cell records below
                rows = struct.unpack_from("<I" if version == 0x0600 else "<H", data, 4 if version == 0x0600 else 2)[0]
                columns = struct.unpack_from("<H", data, 10 if version == 0x0600 else 6)[0]
                limits.check_shape(sheet.name, rows, columns)
            elif code in cell_codes | {0x00BD, 0x00BE}:
                row, first = struct.unpack_from("<HH", data)
                last = struct.unpack_from("<H", data, len(data) - 2)[0] if code in {0x00BD, 0x00BE} else first
                if last < first:
                    raise DatasetError("Некорректный диапазон ячеек XLS")
                sheet.grow(row + 1, last + 1, limits)
                sheet.physical_cells += last - first + 1
                if sheet.physical_cells > limits.cells + limits.columns:
                    raise DatasetError("Лист превышает бюджет физических ячеек")
                if row and code not in {0x0201, 0x00BE}:
                    sheet.nonempty += last - first + 1
            elif code in {0x01B8, 0x0800}:  # HLINK / QUICKTIP may expand cell ranges
                start = 2 if code == 0x0800 else 0
                first_row, last_row, first_col, last_col = struct.unpack_from("<HHHH", data, start)
                limits.check_shape(sheet.name, last_row + 1, last_col + 1)
                if (last_row - first_row + 1) * (last_col - first_col + 1) > limits.cells:
                    raise DatasetError("Диапазон ссылок XLS превышает бюджет ячеек")
        budget.add(sheet)
        sheets.append(sheet)
    return sheets


def _xls_value(cell, datemode):
    if cell.ctype in {xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR}:
        return None
    if cell.ctype == xlrd.XL_CELL_DATE:
        return xlrd.xldate.xldate_as_datetime(cell.value, datemode)
    if cell.ctype == xlrd.XL_CELL_BOOLEAN:
        return bool(cell.value)
    if cell.ctype == xlrd.XL_CELL_NUMBER and cell.value.is_integer():
        return int(cell.value)
    return cell.value


@contextmanager
def read_workbook(content: bytes, suffix: str, limits: ExcelLimits):
    if zipfile.is_zipfile(BytesIO(content)):
        validate_xlsx_archive(content, limits)
        if suffix != ".xlsx":
            raise DatasetError("Содержимое XLSX не соответствует расширению файла .xls")
        bounds = _xlsx_bounds(content, limits)
        workbook = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=True, keep_links=False)
        try:
            if workbook.sheetnames != [sheet.name for sheet in bounds]:
                raise DatasetError("Несогласованный список листов XLSX")

            def sheets():
                for bound in bounds:
                    if not bound.columns:
                        continue
                    sheet = workbook[bound.name]
                    rows = sheet.iter_rows(min_row=1, max_row=bound.rows, min_col=1, max_col=bound.columns)
                    try:
                        yield bound.name, (tuple(None if cell.data_type == "e" else cell.value for cell in row) for row in rows)
                    finally:
                        rows.close()

            with closing(sheets()) as iterator:
                yield iterator
        finally:
            workbook.close()
    elif content.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
        if suffix != ".xls":
            raise DatasetError("Содержимое XLS не соответствует расширению файла .xlsx")
        bounds = _xls_bounds(content, limits)
        with xlrd.open_workbook(file_contents=content, on_demand=True, ragged_rows=True, logfile=StringIO()) as workbook:
            if workbook.nsheets != len(bounds):
                raise DatasetError("Несогласованный список листов XLS")

            def sheets():
                for index, bound in enumerate(bounds):
                    if not bound.columns:
                        continue
                    sheet = workbook.sheet_by_index(index)
                    try:
                        limits.check_shape(sheet.name, sheet.nrows, sheet.ncols)
                        yield sheet.name, (
                            tuple(_xls_value(cell, workbook.datemode) for cell in sheet.row(row))
                            + (None,) * (sheet.ncols - sheet.row_len(row))
                            for row in range(sheet.nrows)
                        )
                    finally:
                        workbook.unload_sheet(index)

            with closing(sheets()) as iterator:
                yield iterator
    else:
        raise DatasetError("Содержимое файла не является XLSX или бинарным XLS")
