from __future__ import annotations

import asyncio
import hashlib
import json
import re
import shutil
import sqlite3
import subprocess
import sys
import zipfile
from concurrent.futures import ThreadPoolExecutor
from contextlib import contextmanager
from io import BytesIO
from pathlib import Path

import openpyxl
import pandas as pd
import pytest
from fastapi.testclient import TestClient
from starlette.datastructures import UploadFile

from app.main import create_app
from app.services import dataset_store as store_module, excel_reader
from app.services.dataset_store import DatasetError, DatasetStore
from app import uploads


@pytest.fixture
def client(tmp_path):
    with TestClient(create_app(tmp_path / "runtime"), raise_server_exceptions=False) as client:
        yield client


def workbook(rows=None, *, extra_sheets=None):
    book = openpyxl.Workbook()
    book.active.title = "Data"
    for row in rows or [["A"], ["value"]]:
        book.active.append(row)
    for name, sheet_rows in (extra_sheets or {}).items():
        sheet = book.create_sheet(name)
        for row in sheet_rows:
            sheet.append(row)
    output = BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def rewrite_xlsx(content, name, transform):
    output = BytesIO()
    with zipfile.ZipFile(BytesIO(content)) as source, zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as dest:
        for item in source.infolist():
            data = source.read(item.filename)
            dest.writestr(item.filename, transform(data) if item.filename == name else data)
    return output.getvalue()


def snapshot(store):
    with store._connect() as connection:
        tables = [list(map(tuple, connection.execute(f"SELECT * FROM {table} ORDER BY 1"))) for table in ("datasets", "records")]
    files = {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in store.upload_dir.iterdir()}
    assert not list(store.runtime_dir.glob(".import-*"))
    return tables, files


def post(client, content, name="book.xlsx"):
    return client.post("/api/datasets", files={"file": (name, content)})


@pytest.mark.parametrize("name,status", [("same.xlsx", 201), ("same.xls", 400)])
def test_detects_xlsx_content_and_rejects_extension_mismatch(client, name, status):
    response = post(client, workbook(), name)
    assert response.status_code == status, response.text
    if status == 400:
        assert "не соответствует" in response.json()["detail"]
        assert snapshot(client.app.state.dataset_store) == ([[], []], {})


@pytest.mark.parametrize("name", ["bomb.xlsx", "bomb.xls"])
def test_zip_bomb_rejected_before_excel_reader_for_both_extensions(client, monkeypatch, name):
    content = BytesIO(workbook())
    # Use the real 100 MiB limit without holding the expanded member in memory.
    with zipfile.ZipFile(content, "a", zipfile.ZIP_DEFLATED) as archive:
        with archive.open("padding.bin", "w") as member:
            for _ in range(101):
                member.write(b"0" * (1024 * 1024))
    monkeypatch.setattr(openpyxl, "load_workbook", lambda *a, **k: pytest.fail("Excel reader ran before ZIP preflight"))
    response = post(client, content.getvalue(), name)
    assert response.status_code == 400
    assert "Распакованный" in response.json()["detail"]
    assert snapshot(client.app.state.dataset_store) == ([[], []], {})


@pytest.mark.parametrize("limit,value,rows,extras", [
    ("MAX_SHEETS", 1, [["A"], ["x"]], {"Second": [["B"], ["y"]]}),
    ("MAX_ROWS_PER_SHEET", 2, [["A"], [1], [2], [3]], {}),
    ("MAX_COLUMNS_PER_SHEET", 2, [["A", "B", "C"], [1, 2, 3]], {}),
    ("MAX_CELLS_PER_SHEET", 5, [["A", "B"], [1, 2], [3, 4], [5, 6]], {}),
    ("MAX_SPARSE_RATIO", 2, [["A", "B", "C"], ["x", None, None]], {}),
    ("MAX_ROWS_PER_WORKBOOK", 3, [["A"], [1], [2]], {"Second": [["B"], [3], [4]]}),
    ("MAX_CELLS_PER_WORKBOOK", 5, [["A", "B"], [1, 2]], {"Second": [["C", "D"], [3, 4], [5, 6]]}),
])
def test_limits_run_before_materialization(client, monkeypatch, limit, value, rows, extras):
    content = workbook(rows, extra_sheets=extras)
    monkeypatch.setattr(store_module, limit, value)
    monkeypatch.setattr(openpyxl, "load_workbook", lambda *a, **k: pytest.fail("Dangerous workbook materialized"))
    monkeypatch.setattr(pd, "read_excel", lambda *a, **k: pytest.fail("Unbounded pandas reader used"))
    response = post(client, content)
    assert response.status_code == 400, response.text
    assert snapshot(client.app.state.dataset_store) == ([[], []], {})


@pytest.mark.parametrize("dimension", [b'<dimension ref="A1:A1"/>', b''])
@pytest.mark.parametrize("coordinate", [b"XFD2", b"A1048576"])
def test_dishonest_or_missing_dimensions_cannot_hide_extreme_cells(client, monkeypatch, dimension, coordinate):
    def transform(data):
        data = re.sub(rb'<dimension[^>]*/>', dimension, data)
        row = re.search(rb"[0-9]+", coordinate)[0]
        return data.replace(b'<row r="2">', b'<row r="' + row + b'">').replace(b'r="A2"', b'r="' + coordinate + b'"')
    content = rewrite_xlsx(workbook(), "xl/worksheets/sheet1.xml", transform)
    monkeypatch.setattr(openpyxl, "load_workbook", lambda *a, **k: pytest.fail("Extreme coordinates reached openpyxl"))
    assert post(client, content).status_code == 400


def test_undersized_dimension_does_not_truncate_valid_data(client):
    content = rewrite_xlsx(workbook([["A", "B"], ["00123", "00456"]]), "xl/worksheets/sheet1.xml",
                           lambda data: re.sub(rb'<dimension[^>]*/>', b'<dimension ref="A1:A1"/>', data))
    response = post(client, content)
    assert response.status_code == 201
    page = client.get(f"/api/datasets/{response.json()['dataset_id']}/records").json()
    assert page["items"][0]["B"] == "00456"


def test_metadata_budget_and_entities_checked_before_openpyxl(client, monkeypatch):
    limits = store_module._excel_limits()
    from dataclasses import replace
    monkeypatch.setattr(store_module, "_excel_limits", lambda: replace(limits, metadata_bytes=64))
    monkeypatch.setattr(openpyxl, "load_workbook", lambda *a, **k: pytest.fail("Unbounded metadata loaded"))
    assert post(client, workbook()).status_code == 400
    monkeypatch.setattr(store_module, "_excel_limits", lambda: limits)
    content = rewrite_xlsx(workbook(), "xl/worksheets/sheet1.xml", lambda data: b'<!DOCTYPE worksheet [<!ENTITY x "unsafe">]>' + data)
    response = post(client, content)
    assert response.status_code == 400
    assert "DTD" in response.json()["detail"]


def test_header_only_width_limit_in_validator_and_real_workbook(client):
    columns = [f"C{i}" for i in range(store_module.MAX_COLUMNS_PER_SHEET + 1)]
    with pytest.raises(DatasetError, match="размер"):
        DatasetStore._validate_frame("Headers", pd.DataFrame(columns=columns))
    assert post(client, workbook([columns])).status_code == 400
    assert post(client, workbook(extra_sheets={"TooWide": [columns]})).status_code == 400


@pytest.mark.parametrize("empty_sheets", [{"Empty": []}, {"Empty": [], "Headers": [["A"]]}])
def test_skips_empty_extra_sheets_and_selects_largest(client, empty_sheets):
    response = post(client, workbook(extra_sheets={**empty_sheets, "Largest": [["A"], ["x"], ["y"]]}))
    assert response.status_code == 201
    assert response.json()["active_sheet"] == "Largest"
    assert [sheet["name"] for sheet in response.json()["sheets"]] == ["Data", "Largest"]


def test_all_empty_workbook_is_rejected(client):
    book = openpyxl.Workbook()
    book.create_sheet("Empty2")
    content = BytesIO()
    book.save(content)
    assert post(client, content.getvalue()).status_code == 400
    assert snapshot(client.app.state.dataset_store) == ([[], []], {})


def test_literal_strings_empty_cells_custom_columns_and_search_values_only(client):
    response = post(client, workbook([
        ["Код", "NA_value", "NULL_value", "Empty", "_customer", "*custom", "Number", "Boolean"],
        ["00123", "NA", "NULL", None, "Иван", "звезда", 0, False],
        ["00456", "NA", "NULL", None, "Анна", "has_value", 2, True],
    ]))
    assert response.status_code == 201, response.text
    url = f"/api/datasets/{response.json()['dataset_id']}/records"
    items = client.get(url).json()["items"]
    assert [item["Код"] for item in items] == ["00123", "00456"]
    assert all(item["NA_value"] == "NA" and item["NULL_value"] == "NULL" and item["Empty"] is None for item in items)
    assert items[0]["_customer"] == "Иван" and items[0]["*custom"] == "звезда"
    for query, expected in [("_", 1), ("_customer", 0), ("_sheet", 0), ("*custom", 0), ('":', 0), ("False", 1), ("0", 2), ("иван", 1), ("%", 0)]:
        assert client.get(url, params={"q": query}).json()["total"] == expected, query


def test_underscore_returns_zero_without_user_underscores(client):
    response = post(client, workbook([["_customer"], ["Иван"], ["Анна"]]))
    assert response.status_code == 201
    assert client.get(f"/api/datasets/{response.json()['dataset_id']}/records", params={"q": "_"}).json()["total"] == 0


@pytest.mark.parametrize("columns", [["A", " A "], [1, "1"], ["A", "A"], ["_sheet"], ["_source_row"]])
def test_original_headers_cannot_be_mangled_to_bypass_validation(client, columns):
    assert post(client, workbook([columns, ["value"] * len(columns)])).status_code == 400


@pytest.mark.parametrize("failure", ["write", "rename", "insert", "commit", "later_sheet", "reader_close"])
def test_failed_import_preserves_existing_rows_and_uploads(client, monkeypatch, failure):
    store = client.app.state.dataset_store
    assert post(client, workbook()).status_code == 201
    before = snapshot(store)
    content = workbook([["A"], [1], [2], [3]], extra_sheets={"Second": [["B"], [4]]})
    if failure == "write":
        original = Path.write_bytes
        def fail(path, data):
            if path.name == "source.xlsx":
                original(path, data[:100])
                raise OSError("disk full")
            return original(path, data)
        monkeypatch.setattr(Path, "write_bytes", fail)
    elif failure == "rename":
        monkeypatch.setattr(Path, "replace", lambda *args: (_ for _ in ()).throw(OSError("rename failed")))
    elif failure == "insert":
        with store._connect() as connection:
            connection.execute("CREATE TRIGGER fail_record BEFORE INSERT ON records WHEN NEW.row_number = 4 BEGIN SELECT RAISE(ABORT, 'insert failed'); END")
    elif failure == "commit":
        original = store._connect
        @contextmanager
        def fail_commit():
            with original() as connection:
                yield connection
                raise sqlite3.OperationalError("commit failed")
        monkeypatch.setattr(store, "_connect", fail_commit)
    elif failure == "reader_close":
        original = store_module.read_workbook
        @contextmanager
        def fail_close(*args, **kwargs):
            with original(*args, **kwargs) as sheets:
                yield sheets
                raise ValueError("reader close failed")
        monkeypatch.setattr(store_module, "read_workbook", fail_close)
    else:
        monkeypatch.setattr(store_module, "INSERT_BATCH_SIZE", 1)
        content = workbook(extra_sheets={"Second": [["B", " B "], [1, 2]]})
    response = post(client, content)
    assert response.status_code in {400, 500}, response.text
    monkeypatch.undo()
    assert snapshot(store) == before


def test_sheets_are_consumed_sequentially_and_inserted_in_bounded_batches(client, monkeypatch):
    original = DatasetStore._insert_sheet
    consumed = []
    def track(self, connection, dataset_id, name, rows):
        if consumed:
            assert consumed[-1][1] == "done"
        consumed.append((name, "reading"))
        result = original(self, connection, dataset_id, name, rows)
        consumed[-1] = (name, "done")
        return result
    monkeypatch.setattr(DatasetStore, "_insert_sheet", track)
    monkeypatch.setattr(pd, "read_excel", lambda *a, **k: pytest.fail("Unbounded DataFrame reader used"))
    monkeypatch.setattr(store_module, "INSERT_BATCH_SIZE", 1)
    assert post(client, workbook(extra_sheets={"Next": [["A"], ["x"]]})).status_code == 201
    assert consumed == [("Data", "done"), ("Next", "done")]


def test_plain_pagination_decodes_only_requested_page(client, monkeypatch):
    response = post(client, workbook([["A"], *[[i] for i in range(21)]]))
    original = json.loads
    decoded_records = []
    def loads(value, *a, **kw):
        if isinstance(value, str) and '"_record_id"' in value:
            decoded_records.append(value)
        return original(value, *a, **kw)
    monkeypatch.setattr(json, "loads", loads)
    page = client.app.state.dataset_store.list_records(response.json()["dataset_id"], limit=1, offset=20)
    assert page.total == 21 and page.items[0]["A"] == 20
    assert len(decoded_records) == 1


def test_search_uses_sqlite_json_values_and_decodes_only_returned_page(client, monkeypatch):
    response = post(client, workbook([["A"], *[[f"row {i}"] for i in range(21)]]))
    original = json.loads
    decoded_records = []

    def loads(value, *args, **kwargs):
        if isinstance(value, str) and '"_record_id"' in value:
            decoded_records.append(value)
        return original(value, *args, **kwargs)

    monkeypatch.setattr(json, "loads", loads)
    page = client.app.state.dataset_store.list_records(
        response.json()["dataset_id"], limit=1, query="row",
    )

    assert page.total == 21 and page.items[0]["A"] == "row 0"
    assert len(decoded_records) == 1


def test_import_main_in_clean_copy_has_no_runtime_side_effects(tmp_path):
    root = Path(__file__).resolve().parents[1]
    shutil.copytree(root / "app", tmp_path / "app", ignore=shutil.ignore_patterns("__pycache__"))
    result = subprocess.run([sys.executable, "-B", "-c", "import app.main; assert callable(app.main.create_app)"], cwd=tmp_path, capture_output=True, text=True)
    assert result.returncode == 0, result.stderr
    assert not (tmp_path / "data").exists()


def test_idempotency_survives_store_restart_without_schema_migration(tmp_path):
    runtime = tmp_path / "runtime"
    content = workbook()
    first = DatasetStore(runtime).import_excel(BytesIO(content), "first.xlsx")
    before = snapshot(DatasetStore(runtime))
    second = DatasetStore(runtime).import_excel(BytesIO(content), "renamed.xlsx")

    assert second.dataset_id == first.dataset_id
    assert second.filename == "first.xlsx"
    assert snapshot(DatasetStore(runtime)) == before


def test_concurrent_identical_imports_create_one_dataset(tmp_path):
    store = DatasetStore(tmp_path / "runtime")
    content = workbook()
    with ThreadPoolExecutor(max_workers=2) as executor:
        results = list(executor.map(
            lambda name: store.import_excel(BytesIO(content), name),
            ["first.xlsx", "second.xlsx"],
        ))

    assert results[0].dataset_id == results[1].dataset_id
    with store._connect() as connection:
        assert connection.execute("SELECT COUNT(*) FROM datasets").fetchone()[0] == 1
        assert connection.execute("SELECT COUNT(*) FROM records").fetchone()[0] == 1
    assert len(list(store.upload_dir.iterdir())) == 1


def multipart_parts():
    return (b'--test-boundary\r\nContent-Disposition: form-data; name="file"; filename="book.xlsx"\r\nContent-Type: application/octet-stream\r\n\r\n', b'\r\n--test-boundary--\r\n')


def asgi_post(application, chunks, headers=()):
    sent, received = [], []
    iterator = iter(chunks)
    async def receive():
        chunk = next(iterator, None)
        if chunk is None:
            return {"type": "http.request", "body": b"", "more_body": False}
        received.append(len(chunk))
        return {"type": "http.request", "body": chunk, "more_body": True}
    async def send(message):
        sent.append(message)
    scope = {"type": "http", "asgi": {"version": "3.0"}, "http_version": "1.1", "method": "POST", "scheme": "http", "path": "/api/datasets", "raw_path": b"/api/datasets", "query_string": b"", "root_path": "", "headers": [(b"content-type", b"multipart/form-data; boundary=test-boundary"), *headers], "client": ("127.0.0.1", 1), "server": ("127.0.0.1", 80)}
    asyncio.run(application(scope, receive, send))
    status = next(item["status"] for item in sent if item["type"] == "http.response.start")
    return status, sum(received)


@pytest.mark.parametrize("headers", [[], [(b"transfer-encoding", b"chunked")], [(b"content-length", b"1")]])
def test_multipart_stops_before_writing_oversized_body(client, monkeypatch, headers):
    written = []
    handles = []
    original = UploadFile.write
    async def track(self, data):
        written.append(len(data))
        handles.append(self.file)
        await original(self, data)
    monkeypatch.setattr(UploadFile, "write", track)
    prefix, suffix = multipart_parts()
    def chunks():
        yield prefix
        for _ in range(26):
            yield b"x" * (1024 * 1024)
        yield suffix
        pytest.fail("The server consumed the oversized stream to completion")
    status, received = asgi_post(client.app, chunks(), headers)
    assert status == 400
    assert sum(written) <= store_module.MAX_UPLOAD_BYTES
    assert received <= store_module.MAX_UPLOAD_BYTES + 1024 * 1024 + len(prefix)
    assert handles and all(handle.closed for handle in handles)
    assert snapshot(client.app.state.dataset_store) == ([[], []], {})


def test_individual_file_limit_is_enforced_before_write_even_with_body_headroom(client, monkeypatch):
    monkeypatch.setattr(uploads, "MAX_UPLOAD_BYTES", 1024)
    written = []
    original = UploadFile.write
    async def track(self, data):
        written.append(len(data))
        await original(self, data)
    monkeypatch.setattr(UploadFile, "write", track)
    prefix, suffix = multipart_parts()
    status, _ = asgi_post(client.app, [prefix, b"x" * 1024, b"x", suffix])
    assert status == 400 and sum(written) <= 1024


def test_small_chunked_multipart_and_missing_content_length_succeed(client):
    prefix, suffix = multipart_parts()
    content = workbook()
    chunks = [prefix, *(content[i:i+97] for i in range(0, len(content), 97)), suffix]
    status, _ = asgi_post(client.app, chunks, [(b"transfer-encoding", b"chunked")])
    assert status == 201


def test_truncated_multipart_closes_temporary_files(client, monkeypatch):
    handles = []
    original = UploadFile.write
    async def track(self, data):
        handles.append(self.file)
        await original(self, data)
    monkeypatch.setattr(UploadFile, "write", track)
    prefix, _ = multipart_parts()
    status, _ = asgi_post(client.app, [prefix, b"partial file"])
    assert status == 400
    assert handles and all(handle.closed for handle in handles)
    assert snapshot(client.app.state.dataset_store) == ([[], []], {})


def test_real_binary_xls_preserves_values_and_skips_empty_sheet(client):
    content = (Path(__file__).parent / "fixtures/values.xls").read_bytes()
    assert content.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1")
    response = post(client, content, "values.xls")
    assert response.status_code == 201, response.text
    dataset = response.json()
    assert dataset["active_sheet"] == "Data"
    assert [sheet["name"] for sheet in dataset["sheets"]] == ["Data", "Second"]
    items = client.get(f"/api/datasets/{dataset['dataset_id']}/records").json()["items"]
    assert [item["Код"] for item in items] == ["00123", "00456"]
    assert all(item["NA"] == "NA" and item["NULL"] == "NULL" and item["Empty"] is None for item in items)
    assert items[0]["Number"] == 0 and items[0]["Boolean"] is False
    assert post(client, content, "values.xlsx").status_code == 400


@pytest.mark.parametrize("limit,value", [
    ("MAX_SHEETS", 2), ("MAX_ROWS_PER_SHEET", 1), ("MAX_COLUMNS_PER_SHEET", 7),
    ("MAX_CELLS_PER_SHEET", 15), ("MAX_ROWS_PER_WORKBOOK", 2),
    ("MAX_CELLS_PER_WORKBOOK", 16), ("MAX_SPARSE_RATIO", 1),
])
def test_binary_xls_limits_checked_before_xlrd_materializes(client, monkeypatch, limit, value):
    content = (Path(__file__).parent / "fixtures/values.xls").read_bytes()
    monkeypatch.setattr(store_module, limit, value)
    monkeypatch.setattr(excel_reader.xlrd, "open_workbook", lambda *a, **k: pytest.fail("XLS reached xlrd before validation"))
    response = post(client, content, "values.xls")
    assert response.status_code == 400, response.text
    assert snapshot(client.app.state.dataset_store) == ([[], []], {})


def test_binary_xls_actual_cell_coordinates_override_dishonest_dimensions(client, monkeypatch):
    import struct
    content = (Path(__file__).parent / "fixtures/values.xls").read_bytes()
    document = excel_reader.CompDoc(content)
    stream = document.get_named_stream("Workbook")
    # The synthetic fixture has a contiguous Workbook stream. Change a LABELSST
    # cell while retaining the small, valid DIMENSIONS record.
    start = content.index(stream)
    position = 0
    while position + 4 <= len(stream):
        code, size = struct.unpack_from("<HH", stream, position)
        if code == 0x00FD and struct.unpack_from("<H", stream, position + 4)[0] == 1:
            mutated = bytearray(content)
            struct.pack_into("<HH", mutated, start + position + 4, 65535, 255)
            break
        position += 4 + size
    else:
        pytest.fail("No data cell in BIFF8 fixture")
    monkeypatch.setattr(excel_reader.xlrd, "open_workbook", lambda *a, **k: pytest.fail("Extreme XLS cell reached xlrd"))
    assert post(client, bytes(mutated), "values.xls").status_code == 400
