"""Read-only Excel audit. Run: python -m postcode.audit --help."""
import argparse
from collections import Counter, defaultdict
from datetime import date, datetime
import hashlib
import json
from pathlib import Path
import re
import unicodedata

from openpyxl import load_workbook


def missing(value):
    return value is None or isinstance(value, str) and not value.strip()


def normalize(value):
    if missing(value):
        return ""
    return re.sub(r"\s+", " ", unicodedata.normalize("NFKC", str(value)).casefold()).strip()


def read_excel(path, sheet):
    book = load_workbook(path, read_only=True, data_only=False)
    try:
        if sheet not in book.sheetnames:
            raise ValueError(f"Нет листа {sheet!r}. Доступны: {book.sheetnames}")
        iterator = book[sheet].iter_rows(values_only=True)
        header = next(iterator, ())
        if not header or all(missing(value) for value in header):
            raise ValueError("Первая строка должна содержать заголовки.")
        names = [str(value).strip() if not missing(value) else "" for value in header]
        if any(not name for name in names) or len(names) != len(set(names)):
            raise ValueError("Заголовки должны быть непустыми и уникальными.")
        rows = [(number, dict(zip(names, values)))
                for number, values in enumerate(iterator, 2)
                if any(not missing(value) for value in values)]
        if not rows:
            raise ValueError("Лист содержит заголовки, но не содержит записей.")
        return names, rows
    finally:
        book.close()


def profile(names, rows, category=None, text_columns=()):
    for name in ([category] if category else []) + list(text_columns):
        if name not in names:
            raise ValueError(f"Не найден столбец {name!r}. Доступны: {names}")
    columns = {}
    for name in names:
        values = [record[name] for _, record in rows]
        present = [value for value in values if not missing(value)]
        dates = [value.isoformat() for value in present if isinstance(value, (date, datetime))]
        columns[name] = {
            "missing": len(values) - len(present),
            "unique": len({(type(value).__name__, str(value)) for value in present}),
            "types": dict(Counter(type(value).__name__ for value in present)),
            "formula_cells": sum(isinstance(value, str) and value.startswith("=") for value in present),
            "date_min": min(dates) if dates else None,
            "date_max": max(dates) if dates else None,
        }
    exact = defaultdict(list)
    texts = defaultdict(list)
    for number, record in rows:
        key = tuple((type(record[name]).__name__, str(record[name])) for name in names)
        exact[key].append(number)
        if text_columns:
            parts = [normalize(record[name]) for name in text_columns]
            if any(parts):
                texts[tuple(parts)].append(number)
    counts = Counter(str(record[category]).strip() for _, record in rows
                     if category and not missing(record[category]))
    ranked = sorted(counts.items(), key=lambda item: (-item[1], item[0]))
    return {
        "row_count": len(rows), "column_count": len(names), "columns": columns,
        "category_column": category, "text_columns": list(text_columns),
        "category_count": len(counts) if category else None,
        "top_15": [{"category": name, "count": count} for name, count in ranked[:15]],
        "outside_top_15": sum(count for _, count in ranked[15:]) if category else None,
        "exact_duplicate_groups": [group for group in exact.values() if len(group) > 1],
        "normalized_text_duplicate_groups": [group for group in texts.values() if len(group) > 1],
        "limitations": [
            "Нормализация текста учитывает регистр, Unicode и пробелы. Поиск близких перефразировок ещё не реализован.",
            "Текстовые даты не преобразуются автоматически. Семантика столбцов и SLA требует проверки.",
            "Формулы не вычисляются. Их наличие отражено в статистике столбцов.",
            "Выборка ещё не разделена. Отчёт не является оценкой качества модели.",
        ],
    }


def markdown(report):
    lines = ["# Аудит данных", "", f"Записей: {report['row_count']}. Столбцов: {report['column_count']}.",
             "", f"SHA-256: `{report['sha256']}`", "", "## Столбцы", ""]
    for name, values in report["columns"].items():
        lines.append(f"- {name}: пропусков {values['missing']}, уникальных значений {values['unique']}, формул {values['formula_cells']}.")
    lines += ["", "## 15 наиболее частых видов", ""]
    if report["category_column"]:
        lines += [f"- {item['category']}: {item['count']}" for item in report["top_15"]]
    else:
        lines.append("Столбец вида не указан. Повторите запуск с --category.")
    lines += ["", "## Дубликаты", "",
              f"Групп полных дубликатов: {len(report['exact_duplicate_groups'])}.",
              f"Групп совпадающих нормализованных текстов: {len(report['normalized_text_duplicate_groups'])}.",
              "Номера строк Excel указаны в JSON. Тексты обращений в отчёт не включаются."]
    if not report["text_columns"]:
        lines.append("Текстовые столбцы не заданы: проверка совпадений текстов не выполнялась.")
    lines += ["", "## Ограничения", ""] + [f"- {item}" for item in report["limitations"]]
    return "\n".join(lines) + "\n"


def main():
    parser = argparse.ArgumentParser(description="Аудит Excel без изменения исходного файла")
    parser.add_argument("input", type=Path)
    parser.add_argument("--sheet", default="Sheet0")
    parser.add_argument("--category", help="Точное имя столбца вида обращения")
    parser.add_argument("--text", action="append", default=[], help="Столбец текста при регистрации; можно повторять")
    parser.add_argument("--output", type=Path, default=Path("reports/audit"))
    args = parser.parse_args()
    try:
        names, rows = read_excel(args.input, args.sheet)
        report = profile(names, rows, args.category, args.text)
        report.update(sha256=hashlib.sha256(args.input.read_bytes()).hexdigest(), sheet=args.sheet)
        json_path = args.output.with_suffix(".json")
        md_path = args.output.with_suffix(".md")
        if args.input.resolve() in (json_path.resolve(), md_path.resolve()):
            raise ValueError("Путь отчёта совпадает с исходным файлом.")
        args.output.parent.mkdir(parents=True, exist_ok=True)
        json_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        md_path.write_text(markdown(report), encoding="utf-8")
        print(f"Готово: {json_path}, {md_path}")
    except (OSError, ValueError) as error:
        parser.exit(2, f"Ошибка: {error}\n")


if __name__ == "__main__":
    main()
